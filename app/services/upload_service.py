import csv
import io
import json
import re
from typing import Any, List, Tuple
from pydantic import ValidationError

from app.schemas.arena import QuestionSchema


def _validate_and_coerce_schema(obj: dict, idx: int, errors: List[dict]) -> QuestionSchema | None:
    """Helper to handle common data normalization and validate against QuestionSchema."""
    try:
        # Normalize options if present as string
        if "options" in obj and not isinstance(obj["options"], list):
            obj["options"] = _parse_options_cell(obj.get("options"))

        # Coerce boolean-like values for AI flag
        if "is_ai_generated" in obj:
            v = obj["is_ai_generated"]
            if isinstance(v, str):
                obj["is_ai_generated"] = v.lower() in ("1", "true", "yes")

        return QuestionSchema.model_validate(obj)
    except ValidationError as ve:
        errors.append({"row": idx, "messages": [e["msg"] for e in ve.errors()]})
    except Exception as e:
        errors.append({"row": idx, "messages": [str(e)]})
    return None


def _parse_options_cell(raw: Any) -> List[str]:
    """Normalize an options cell into a list of strings.
    Accepts JSON arrays, pipe-separated or semicolon-separated values, or a Python list.
    """
    if raw is None:
        return []

    if isinstance(raw, list):
        return [str(x).strip() for x in raw if str(x).strip()]

    s = str(raw).strip()
    
    # Try JSON array format first
    if s.startswith("[") and s.endswith("]"):
        try:
            arr = json.loads(s)
            if isinstance(arr, list):
                return [str(x).strip() for x in arr if str(x).strip()]
        except (json.JSONDecodeError, TypeError):
            pass

    # Fallback character splitters
    if "|" in s:
        return [p.strip() for p in s.split("|") if p.strip()]
    if ";" in s:
        return [p.strip() for p in s.split(";") if p.strip()]

    # Fallback to single string option
    return [s] if s else []


def _map_tabular_row(row: dict) -> dict:
    """Maps unstructured dictionary keys from CSV/Excel formats into standard schema keys."""
    raw_opts = row.get("options") or row.get("choices") or row.get("answers")
    
    return {
        "prompt_text": row.get("prompt_text") or row.get("prompt") or row.get("question"),
        "time_limit_seconds": int(row.get("time_limit_seconds") or row.get("time_limit") or 10),
        "point_value": int(row.get("point_value") or 10),
        "correct_option_index": int(row.get("correct_option_index") or 0),
        "status": row.get("status") or "ready",
        "is_ai_generated": row.get("is_ai_generated"),
        "options": _parse_options_cell(raw_opts),
    }


def _parse_xlsx(content_bytes: bytes) -> List[dict]:
    """Return list of row dicts from first sheet of an XLSX workbook."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is not installed")

    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
    ws = wb.active
    
    if ws is None:
        return []
        
    rows = list(ws.iter_rows(values_only=True))

    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    results = []
    for row in rows[1:]:
        rowdict = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            rowdict[h] = row[i] if i < len(row) else None
        results.append(rowdict)
    return results


def _extract_text_from_docx(content_bytes: bytes) -> str:
    """Extract paragraphs from a .docx file buffer."""
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx is not installed")

    doc = Document(io.BytesIO(content_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n".join(paragraphs)


def _extract_text_from_pdf(content_bytes: bytes) -> str:
    """Extract standard text layers from a .pdf file buffer."""
    try:
        from pdfminer.high_level import extract_text_to_fp
    except ImportError:
        raise RuntimeError("pdfminer.six is not installed")

    out = io.StringIO()
    try:
        extract_text_to_fp(io.BytesIO(content_bytes), out)
        return out.getvalue()
    except Exception as e:
        raise RuntimeError(f"PDF extraction failed: {e}")


def _ocr_pdf_with_tesseract(content_bytes: bytes, lang: str = "eng") -> str:
    """Fallback method: Perform local OCR on standard PDF documents using Tesseract."""
    try:
        from pdf2image import convert_from_bytes
        import pytesseract
    except ImportError as e:
        raise RuntimeError(f"pytesseract/pdf2image not available: {e}")

    try:
        images = convert_from_bytes(content_bytes)
    except Exception as e:
        raise RuntimeError(f"Failed to convert PDF to images for OCR: {e}")

    text_pieces = []
    for img in images:
        try:
            text = pytesseract.image_to_string(img, lang=lang)
            text_pieces.append(text)
        except Exception:
            text_pieces.append("")  # Fallback to blank page on localized errors

    return "\n".join(text_pieces)


def _ocr_with_google_vision(content_bytes: bytes) -> str:
    """Fallback method: Perform cloud OCR using Google Vision APIs."""
    try:
        from google.cloud import vision
        from pdf2image import convert_from_bytes
    except ImportError as e:
        raise RuntimeError(f"google-cloud-vision or pdf2image not available: {e}")

    client = vision.ImageAnnotatorClient()
    images = convert_from_bytes(content_bytes)
    text_pieces = []
    
    for img in images:
        try:
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            content = buf.getvalue()
            image = vision.Image(content=content)
            response = client.text_detection(image=image)
            
            if response.error.message:
                text_pieces.append("")
            else:
                desc = response.text_annotations[0].description if response.text_annotations else ""
                text_pieces.append(desc)
        except Exception:
            text_pieces.append("")

    return "\n".join(text_pieces)


def _parse_unstructured_text(text: str) -> List[dict]:
    """Attempt to split raw unstructured text into structural question definitions."""
    text = re.sub(r"\r\n?", "\n", text)
    splits = re.split(r"\n(?=\s*(?:Q?\d{1,3}[\)\.]|Question\s+\d+))", text, flags=re.IGNORECASE)
    results = []
    
    for block in splits:
        b = block.strip()
        if not b:
            continue

        lines = [ln.strip() for ln in b.split("\n") if ln.strip()]
        if not lines:
            continue

        options = []
        prompt_lines = []
        for ln in lines:
            m = re.match(r"^[\(]?([A-D])\)?[\.|\)]\s*(.+)$", ln, flags=re.IGNORECASE)
            if m:
                options.append(m.group(2).strip())
            else:
                if re.match(r"^[\-\u2022]\s+", ln):
                    options.append(re.sub(r"^[\-\u2022]\s+", "", ln).strip())
                else:
                    prompt_lines.append(ln)

        if not options:
            inline_opts = re.findall(r"([A-D][\)\.]\s*[^A-D\)\.]*)", b)
            if inline_opts:
                options = [re.sub(r"^[A-D][\)\.]\s*", "", s).strip() for s in inline_opts]

        prompt_text = " ".join(prompt_lines).strip() if prompt_lines else (options and options.pop(0)) or b

        correct_index = 0
        ans_match = re.search(r"Answer[:\s]+([A-D]|\d+)", b, flags=re.IGNORECASE)
        if ans_match:
            val = ans_match.group(1)
            if val.isdigit():
                try:
                    correct_index = max(0, int(val) - 1)
                except ValueError:
                    correct_index = 0
            else:
                correct_index = ord(val.upper()) - 65

        results.append({
            "prompt_text": prompt_text,
            "time_limit_seconds": 10,
            "options": options,
            "correct_option_index": correct_index,
            "point_value": 10,
        })

    return results


def parse_questions_file(content_bytes: bytes, filename: str = "upload") -> Tuple[List[QuestionSchema], List[dict]]:
    """Parse an uploaded questions file contextually based on extension.

    Returns a tuple (valid_questions, errors) where errors is a list of {row, messages}.
    """
    # Decoding textual file layers safely
    text = None
    try:
        text = content_bytes.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = content_bytes.decode("latin-1")
        except UnicodeDecodeError:
            return [], [{"row": 0, "messages": ["Failed to decode file (not UTF-8 or latin-1)"]}]

    ext = filename.lower()
    results: List[QuestionSchema] = []
    errors: List[dict] = []

    # Process JSON Layout
    if ext.endswith(".json"):
        try:
            arr = json.loads(text)
            if not isinstance(arr, list):
                return [], [{"row": 0, "messages": ["JSON root must be an array of question objects"]}]

            for idx, obj in enumerate(arr, start=1):
                q = _validate_and_coerce_schema(obj, idx, errors)
                if q:
                    results.append(q)
        except json.JSONDecodeError as e:
            return [], [{"row": 0, "messages": [f"Invalid JSON: {str(e)}"]}]

    # Process XLSX Layout
    elif ext.endswith(".xlsx") or ext.endswith(".xls"):
        try:
            rows = _parse_xlsx(content_bytes)
            for idx, row in enumerate(rows, start=1):
                mapped_obj = _map_tabular_row(row)
                q = _validate_and_coerce_schema(mapped_obj, idx, errors)
                if q:
                    results.append(q)
        except Exception as e:
            return [], [{"row": 0, "messages": [f"Failed to parse XLSX: {str(e)}"]}]

    # Process Document (DOCX/PDF) Layouts
    elif ext.endswith(".docx") or ext.endswith(".pdf"):
        try:
            if ext.endswith(".docx"):
                text = _extract_text_from_docx(content_bytes)
            else:
                try:
                    text = _extract_text_from_pdf(content_bytes)
                except Exception:
                    text = ""

                # OCR Processing Fallback
                if not text or len(text.strip()) < 80:
                    try:
                        text = _ocr_pdf_with_tesseract(content_bytes)
                    except Exception:
                        try:
                            text = _ocr_with_google_vision(content_bytes)
                        except Exception as e:
                            raise RuntimeError(f"OCR failed: {e}")

            candidates = _parse_unstructured_text(text)
            for idx, obj in enumerate(candidates, start=1):
                q = _validate_and_coerce_schema(obj, idx, errors)
                if q:
                    results.append(q)
        except Exception as e:
            return [], [{"row": 0, "messages": [f"Failed to extract text/OCR: {str(e)}"]}]

    # Process Standard CSV Layout (Default Fallback)
    else:
        try:
            decoded = io.StringIO(text)
            reader = csv.DictReader(decoded)
            for idx, row in enumerate(reader, start=1):
                mapped_obj = _map_tabular_row(row)
                q = _validate_and_coerce_schema(mapped_obj, idx, errors)
                if q:
                    results.append(q)
        except Exception as e:
            return [], [{"row": 0, "messages": [f"Failed to parse CSV: {str(e)}"]}]

    return results, errors


def ai_parse_text_to_questions(text: str, api_key: str) -> Tuple[List[QuestionSchema], List[dict]]:
    """Use Gemini (via google.genai) to parse raw text into a valid QuestionSchema array."""
    try:
        from google.genai import Client
    except ImportError as e:
        return [], [{"row": 0, "messages": [f"AI client package not available: {e}"]}]

    client = Client(api_key=api_key)
    prompt = [
        {
            "role": "user",
            "content": (
                "Extract all quiz questions from the text below and return ONLY a JSON array of objects. "
                "Each object must include: prompt_text (string), options (array of strings), correct_option_index (0-based integer), "
                "time_limit_seconds (integer, default 10), point_value (integer, default 10). "
                "If an answer is not clear, set correct_option_index to 0. Do not include any extra text.\n\n"
                "TEXT:\n" + text
            ),
        }
    ]

    try:
        response = client.models.generate_content(
            model="gemini-3.5-flash", 
            contents=prompt, 
            config={"temperature": 0.0}
        )
        resp_text = response.text or ""
        
        # Clean potential markdown wrappers safely
        resp_text = resp_text.replace("```json", "").replace("```", "").strip()
        try:
            arr = json.loads(resp_text)
        except json.JSONDecodeError as e:
            return [], [{"row": 0, "messages": [f"AI returned invalid JSON structure: {e}"]}]

        questions: List[QuestionSchema] = []
        errors: List[dict] = []
        
        for idx, obj in enumerate(arr, start=1):
            # Enforce pipeline schemas
            obj.setdefault("time_limit_seconds", 10)
            obj.setdefault("point_value", 10)
            obj.setdefault("correct_option_index", 0)
            obj.setdefault("is_ai_generated", True)

            q = _validate_and_coerce_schema(obj, idx, errors)
            if q:
                q.is_ai_generated = True
                questions.append(q)

        return questions, errors

    except Exception as e:
        return [], [{"row": 0, "messages": [f"AI parsing engine failure: {e}"]}]